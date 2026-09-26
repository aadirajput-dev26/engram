"""
DOCX parser using python-docx.

Paragraphs and native tables are extracted from the .docx file.
The full document is treated as a single ParsedPage since python-docx
does not expose a reliable page-boundary API (pagination is a render-time
concept in Word). Text is returned as a flat string; native tables are
converted to DetectedTable objects.
"""
from __future__ import annotations

from app.core.logging import get_logger
from app.core.parsing.pdf_parser import DetectedTable, ParsedDocument, ParsedPage

logger = get_logger(__name__)


def parse_docx(file_path: str) -> ParsedDocument:
    """
    Parse a .docx file using python-docx.

    All paragraphs are joined into raw_text.
    All tables in the document are extracted as DetectedTable records.
    The result is represented as a single ParsedPage (page 1).

    Args:
        file_path: Absolute path to the .docx file.

    Returns:
        ParsedDocument with a single page containing all extracted content.
    """
    try:
        from docx import Document as DocxDocument  # type: ignore
    except ImportError:
        raise RuntimeError(
            "python-docx is required for DOCX parsing. "
            "Install it with: pip install python-docx"
        )

    try:
        doc = DocxDocument(file_path)
    except Exception as e:
        logger.error("Failed to open DOCX %s: %s", file_path, e)
        raise

    # --- Extract paragraphs ---
    paragraph_lines: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraph_lines.append(text)

    raw_text = "\n".join(paragraph_lines)

    # --- Extract tables ---
    all_tables: list[DetectedTable] = []
    table_texts: list[str] = []

    for table_idx, table in enumerate(doc.tables):
        rows: list[list[str]] = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            rows.append(cells)

        if not rows:
            continue

        header_row = rows[0]
        dt = DetectedTable(
            page_number=1,
            row_count=len(rows),
            col_count=len(header_row),
            header_row=header_row,
            raw_cells=rows,
        )
        all_tables.append(dt)

        # Append markdown representation to the page text
        table_md = _rows_to_markdown(header_row, rows[1:])
        table_texts.append(f"\n[Table {table_idx + 1}]\n{table_md}")

    # Combine paragraphs + table representations
    if table_texts:
        raw_text = raw_text + "\n" + "\n".join(table_texts)

    page = ParsedPage(
        page_number=1,
        raw_text=raw_text,
        source_type="native",
        text_density=len(raw_text),
        tables=all_tables,
        blocks=[],
    )

    logger.info(
        "Parsed DOCX: %d paragraphs, %d tables extracted",
        len(paragraph_lines), len(all_tables),
    )
    return ParsedDocument(
        page_count=1,
        pages=[page],
        detected_type="docx",
        tables=all_tables,
    )


def _rows_to_markdown(header: list[str], rows: list[list[str]]) -> str:
    """Render table rows as a GitHub-style markdown table string."""
    if not header:
        return ""

    col_widths = [max(len(h), 1) for h in header]
    for row in rows:
        for i, cell in enumerate(row):
            if i < len(col_widths):
                col_widths[i] = max(col_widths[i], len(str(cell)))

    def fmt_row(cells: list[str]) -> str:
        parts = []
        for i, cell in enumerate(cells):
            width = col_widths[i] if i < len(col_widths) else len(cell)
            parts.append(str(cell).ljust(width))
        return "| " + " | ".join(parts) + " |"

    separator = "| " + " | ".join("-" * w for w in col_widths) + " |"
    lines = [fmt_row(header), separator]
    for row in rows:
        lines.append(fmt_row(row))
    return "\n".join(lines)
